"""監視対象の一括登録。今お使いのスプレッドシート（.xlsx）や CSV をそのまま読み込める。

必要な列: 販売商品URL（必須）、サイト、店舗、商品名、アラート価格1、アラート価格2、ポイント率（任意）
英語の列名 (url, site, shop, product, alert_price1, alert_price2, point_rate) も可。
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import load_workbook

from .db import Database
from .models import Listing
from .parsers import site_name_for_url
from .parsers.base import normalize, parse_yen

COLUMN_ALIASES = {
    "url": ("販売商品url", "url", "商品url", "販売url"),
    "site": ("サイト", "site", "モール"),
    "shop": ("店舗", "shop", "ショップ", "店舗名"),
    "product": ("商品名", "product", "商品"),
    "alert_price1": ("アラート価格1", "alert_price1", "alert1"),
    "alert_price2": ("アラート価格2", "alert_price2", "alert2"),
    "point_rate": ("ポイント率", "point_rate", "ポイント還元率"),
}


def _match_columns(header: list) -> dict[str, int]:
    norm = [normalize(str(h or "")) for h in header]
    out = {}
    for key, aliases in COLUMN_ALIASES.items():
        for idx, h in enumerate(norm):
            if h in {normalize(a) for a in aliases}:
                out[key] = idx
                break
    return out


def _rows_from_file(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return list(csv.reader(io.StringIO(text)))


def _to_int(v) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    return parse_yen(str(v))


def _to_float(v) -> float | None:
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace("%", "").replace("％", ""))
    except ValueError:
        return None


def parse_listings(rows: list[list]) -> list[Listing]:
    # ヘッダ行を探す（上にハイライト設定行などがあってもよい）
    header_idx, cols = None, {}
    for i, row in enumerate(rows[:20]):
        cols = _match_columns(row)
        if "url" in cols:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("「販売商品URL」(または url) 列が見つかりません")

    def get(row, key):
        idx = cols.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    listings = []
    for row in rows[header_idx + 1:]:
        url = str(get(row, "url") or "").strip()
        if not url.startswith("http"):
            continue
        site = str(get(row, "site") or "").strip() or site_name_for_url(url)
        shop = str(get(row, "shop") or "").strip() or site
        listings.append(
            Listing(
                site=site,
                shop=shop,
                product=str(get(row, "product") or "").strip() or "(未設定)",
                url=url,
                alert_price1=_to_int(get(row, "alert_price1")),
                alert_price2=_to_int(get(row, "alert_price2")),
                point_rate=_to_float(get(row, "point_rate")),
            )
        )
    return listings


def import_file(db: Database, path: str | Path) -> int:
    listings = parse_listings(_rows_from_file(Path(path)))
    for l in listings:
        db.upsert_listing(l)
    return len(listings)
