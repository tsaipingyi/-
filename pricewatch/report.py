"""時系列の価格表（スプレッドシートと同じ「行=店舗、列=取得時刻」）を作り、Excel に書き出す。"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import Database
from .models import Listing

PINK = "FFC0CB"  # 値下げ
BLUE = "87CEEB"  # 値上げ
YELLOW = "FFFF00"  # 指定価格ハイライト
LIGHT_YELLOW = "FFF59D"


@dataclass
class Cell:
    value: int | None
    change: str | None = None  # "down" / "up"
    lowest: bool = False  # 同じ商品・同じ時刻での最安
    below_alert: bool = False
    error: str | None = None
    points: int = 0
    price: int | None = None


@dataclass
class Row:
    listing: Listing
    cells: list[Cell] = field(default_factory=list)

    @property
    def latest(self) -> int | None:
        for c in reversed(self.cells):
            if c.value is not None:
                return c.value
        return None


@dataclass
class Matrix:
    runs: list[dict]
    rows: list[Row]

    @property
    def headers(self) -> list[str]:
        return [header_label(r["started_at"]) for r in self.runs]


def header_label(dt: datetime) -> str:
    return f"{dt:%Y/%m/%d %H:%M}〜"


def build_matrix(db: Database, limit: int = 30, product: str | None = None) -> Matrix:
    runs = db.runs(limit=limit)
    run_ids = [r["id"] for r in runs]
    snaps = db.snapshots_for_runs(run_ids)
    rows: list[Row] = []
    for listing in db.listings(product=product):
        row = Row(listing)
        prev_value = None
        for rid in run_ids:
            s = snaps.get((listing.id, rid))
            value = s["effective"] if s else None
            cell = Cell(value=value, error=s["error"] if s else None,
                        points=s["points"] if s else 0, price=s["price"] if s else None)
            if value is not None:
                if prev_value is not None and value < prev_value:
                    cell.change = "down"
                elif prev_value is not None and value > prev_value:
                    cell.change = "up"
                thresholds = [t for t in (listing.alert_price1, listing.alert_price2) if t is not None]
                cell.below_alert = any(value < t for t in thresholds)
                prev_value = value
            row.cells.append(cell)
        rows.append(row)

    # 商品ごと・時刻ごとの最安値に印
    for i in range(len(run_ids)):
        mins: dict[str, int] = {}
        for row in rows:
            v = row.cells[i].value
            if v is not None:
                mins[row.listing.product] = min(mins.get(row.listing.product, v), v)
        for row in rows:
            c = row.cells[i]
            if c.value is not None and c.value == mins.get(row.listing.product):
                c.lowest = True
    return Matrix(runs=runs, rows=rows)


HEADERS = ["サイト", "店舗", "商品名", "販売商品URL", "アラート価格1", "アラート価格2"]
FIRST_DATA_ROW = 5
FIRST_PRICE_COL = len(HEADERS) + 1


def export_excel(db: Database, path: str | Path, limit: int = 30, product: str | None = None) -> Path:
    m = build_matrix(db, limit=limit, product=product)
    wb = Workbook()
    ws = wb.active
    ws.title = "価格推移"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    # 1〜2行目: ハイライト条件（Excel 上で値を変えると条件付き書式が即反映）
    ws["A1"], ws["C1"] = "価格が", "のセルをハイライト"
    ws["A2"], ws["C2"] = "さらに ±", "のセルをハイライト"
    ws["B1"].fill = PatternFill("solid", fgColor=YELLOW)
    ws["B2"].fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    for ref in ("A1", "A2"):
        ws[ref].alignment = Alignment(horizontal="right")

    # 3行目: 取得時刻 / 4行目: 項目名
    for j, label in enumerate(HEADERS, 1):
        c = ws.cell(row=4, column=j, value=label)
        c.fill, c.border, c.font = header_fill, border, Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for k, label in enumerate(m.headers):
        col = FIRST_PRICE_COL + k
        top = ws.cell(row=3, column=col, value=label)
        top.border, top.alignment = border, Alignment(horizontal="center")
        sub = ws.cell(row=4, column=col, value="商品価格-ポイント")
        sub.fill, sub.border, sub.font = header_fill, border, Font(bold=True)
        sub.alignment = Alignment(horizontal="center")

    pink = PatternFill("solid", fgColor=PINK)
    blue = PatternFill("solid", fgColor=BLUE)
    for i, row in enumerate(m.rows):
        r = FIRST_DATA_ROW + i
        l = row.listing
        for j, v in enumerate([l.site, l.shop, l.product, l.url, l.alert_price1, l.alert_price2], 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = border
            if j >= 5:
                c.number_format = "#,##0"
        ws.cell(row=r, column=4).hyperlink = l.url
        for k, cell in enumerate(row.cells):
            c = ws.cell(row=r, column=FIRST_PRICE_COL + k, value=cell.value)
            c.border = border
            c.number_format = "#,##0"
            if cell.value is None:
                continue
            c.font = Font(bold=True, color="FF0000" if cell.lowest else "000000")
            if cell.change == "down":
                c.fill = pink
            elif cell.change == "up":
                c.fill = blue

    last_row = FIRST_DATA_ROW + max(len(m.rows), 1) - 1
    last_col = FIRST_PRICE_COL + max(len(m.runs), 1) - 1
    if m.runs and m.rows:
        first = f"{get_column_letter(FIRST_PRICE_COL)}{FIRST_DATA_ROW}"
        rng = f"{first}:{get_column_letter(last_col)}{last_row}"
        ws.conditional_formatting.add(
            rng,
            FormulaRule(
                formula=[f'AND(ISNUMBER($B$1),ISNUMBER({first}),ABS({first}-$B$1)<=N($B$2))'],
                fill=PatternFill("solid", fgColor=YELLOW),
                stopIfTrue=False,
            ),
        )

    ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=FIRST_PRICE_COL)
    widths = [14, 22, 18, 24, 12, 12]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for k in range(len(m.runs)):
        ws.column_dimensions[get_column_letter(FIRST_PRICE_COL + k)].width = 17

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
