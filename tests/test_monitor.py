from openpyxl import Workbook, load_workbook

from pricewatch.db import Database
from pricewatch.fetch import FetchError
from pricewatch.importer import import_file
from pricewatch.models import Listing
from pricewatch.monitor import BELOW_ALERT, BACK_IN_STOCK, NEW_LOWEST, OUT_OF_STOCK, PRICE_DOWN, PRICE_UP, run_check
from pricewatch.report import FIRST_DATA_ROW, FIRST_PRICE_COL, build_matrix, export_excel


def page(price):
    return f'<html><head><meta itemprop="price" content="{price}"></head><body></body></html>'


class FakeFetcher:
    """URL → 価格 を差し替えながら巡回をシミュレートする。"""

    use_browser = False
    auto_browser = False

    def __init__(self):
        self.prices = {}
        self.calls = 0

    def clear_cache(self):
        pass

    def get(self, url, browser=None):
        self.calls += 1
        price = self.prices.get(url)
        if price is None:
            raise FetchError("HTTP 503")
        return page(price)


def setup_db():
    db = Database(":memory:")
    a = db.upsert_listing(Listing("楽天市場", "店A", "FIT 4", "https://item.rakuten.co.jp/a/1/", alert_price2=20000))
    b = db.upsert_listing(Listing("ビックカメラ", "ビックカメラ", "FIT 4", "https://www.biccamera.com/bc/item/1/",
                                  point_rate=10))
    return db, a, b


def kinds(events):
    return sorted(e.kind for e in events)


def test_events_over_runs():
    db, a, b = setup_db()
    f = FakeFetcher()
    urls = [l.url for l in db.listings()]

    f.prices = {urls[0]: 21000, urls[1]: 22000}
    _, ev = run_check(db, f)
    assert ev == []  # 初回は比較対象なし

    # ポイント率10% → 22000-2200 = 19800
    snap = db.history(b)[0]
    assert (snap["price"], snap["points"], snap["effective"]) == (22000, 2200, 19800)

    f.prices = {urls[0]: 19500, urls[1]: 22000}
    _, ev = run_check(db, f)
    assert kinds(ev) == sorted([PRICE_DOWN, BELOW_ALERT, NEW_LOWEST])

    # 下回ったままなら再通知しない
    f.prices = {urls[0]: 19400, urls[1]: 22000}
    _, ev = run_check(db, f)
    assert BELOW_ALERT not in kinds(ev)

    f.prices = {urls[1]: 23000}
    _, ev = run_check(db, f)
    assert kinds(ev) == sorted([OUT_OF_STOCK, PRICE_UP])

    f.prices = {urls[0]: 19400, urls[1]: 23000}
    _, ev = run_check(db, f)
    assert kinds(ev) == [BACK_IN_STOCK]


def test_matrix_and_excel(tmp_path):
    db, a, b = setup_db()
    f = FakeFetcher()
    urls = [l.url for l in db.listings()]
    for pa, pb in [(21000, 22000), (19500, 22000), (19500, 24000)]:
        f.prices = {urls[0]: pa, urls[1]: pb}
        run_check(db, f)

    m = build_matrix(db)
    row_a, row_b = m.rows
    assert [c.value for c in row_a.cells] == [21000, 19500, 19500]
    assert [c.change for c in row_a.cells] == [None, "down", None]
    assert [c.change for c in row_b.cells] == [None, None, "up"]
    # 1回目は B(19800) が最安、2回目以降は A
    assert [c.lowest for c in row_a.cells] == [False, True, True]
    assert row_a.cells[1].below_alert

    out = export_excel(db, tmp_path / "out.xlsx")
    ws = load_workbook(out).active
    assert ws.cell(row=4, column=4).value == "販売商品URL"
    assert ws.cell(row=4, column=FIRST_PRICE_COL).value == "商品価格-ポイント"
    assert ws.cell(row=3, column=FIRST_PRICE_COL).value.endswith("〜")
    c = ws.cell(row=FIRST_DATA_ROW, column=FIRST_PRICE_COL + 1)
    assert c.value == 19500 and c.fill.fgColor.rgb.endswith("FFC0CB") and c.font.color.rgb.endswith("FF0000")


def test_kakaku_page_fetched_once_per_run():
    db = Database(":memory:")
    url = "https://kakaku.com/item/K0001/"
    for shop in ("店A", "店B"):
        db.upsert_listing(Listing("価格.com", shop, "FIT 4", url))
    from pricewatch.fetch import Fetcher

    class Session:
        calls = 0

        def get(self, *a, **k):
            Session.calls += 1

            class R:
                status_code = 200
                encoding = "utf-8"
                content = "<table><tr><td>店A</td><td>￥100</td></tr><tr><td>店B</td><td>￥200</td></tr></table>".encode()
            return R()

    _, _ = run_check(db, Fetcher(delay=0, session=Session()))
    assert Session.calls == 1
    assert [db.history(l.id)[0]["price"] for l in db.listings()] == [100, 200]


def test_import_spreadsheet_layout(tmp_path):
    """ユーザーのスプレッドシートと同じ形（上2行が設定、4行目がヘッダ）を読み込めること。"""
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "価格が", "のセルをハイライト"
    ws.append([])
    ws.append([None] * 6 + ["2026/09/29 10:00〜"])
    ws.append(["サイト", "店舗", "商品名", "販売商品URL", "アラート価格1", "アラート価格2", "商品価格-ポイント"])
    ws.append(["Amazon.co.jp", "Amazon.co.jp", "FIT 4 White", "https://www.amazon.co.jp/dp/X", None, 21582, 16414])
    ws.append(["価格.com", "ソフマップ.com", "FIT 4 White", "https://kakaku.com/item/K1/", None, "21,582", 16315])
    ws.append([None, None, None, None, None, None, None])
    path = tmp_path / "sheet.xlsx"
    wb.save(path)

    db = Database(":memory:")
    assert import_file(db, path) == 2
    ls = db.listings()
    assert [(l.site, l.shop, l.alert_price2) for l in ls] == [
        ("Amazon.co.jp", "Amazon.co.jp", 21582),
        ("価格.com", "ソフマップ.com", 21582),
    ]
    assert import_file(db, path) == 2 and len(db.listings()) == 2  # 再取込は上書き
